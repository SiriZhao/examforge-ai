from __future__ import annotations

import time
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.config import settings
from app.main import app
from app.schemas.review import (
    AnkiCard,
    DocumentSection,
    DocumentStructure,
    GenerateReviewResponse,
    LLMConfig,
    LLMErrorInfo,
    ParsedFile,
    ParsedPage,
    ReviewReport,
    StudyUnit,
)
from app.services import generation_store
from app.services.export_service import ExportError, export_anki_csv, export_review_report
from app.services.generation_diagnostics import diagnostic, diagnostic_context
from app.services.generation_pipeline import (
    ChapterSource,
    build_chapter_sources,
    generate_hierarchical_report,
    split_chapter,
)
from app.services.llm_capabilities import capability_for, estimate_tokens
from app.services.llm_providers import get_llm_provider
from app.services.llm_providers.base import LLMProviderError
from app.services.review_planner import generate_review_report


def parsed_material(text: str = "Definition: vector space. Formula: Ax=b.") -> ParsedFile:
    pages = [
        ParsedPage(
            page_number=1,
            text=text,
            source="text_extract",
            source_anchor="linear.pdf, p. 1",
        )
    ]
    return ParsedFile(
        filename="linear.pdf",
        file_type=".pdf",
        path="linear.pdf",
        pages=pages,
        raw_text=text,
        document_structure=DocumentStructure(
            document_id="doc-1",
            filename="linear.pdf",
            file_type=".pdf",
            title="Linear algebra",
            page_count=1,
            sections=[
                DocumentSection(
                    section_id="s1",
                    title="Vector spaces",
                    page_start=1,
                    page_end=1,
                    block_ids=[],
                    source_anchor="linear.pdf, p. 1",
                )
            ],
        ),
    )


def unit(name: str = "Vector spaces") -> StudyUnit:
    return StudyUnit(
        name=name,
        reason="Core chapter",
        priority=90,
        must_know=["Definition"],
        key_points=["Basis", "Dimension"],
        formulas_or_methods=["Ax=b"],
        common_exam_angles=["Proof"],
        pitfalls=["Linear dependence"],
        how_to_review="Recall then solve.",
    )


def test_token_budget_is_model_capability_driven_and_cjk_conservative() -> None:
    capability = capability_for("deepseek", "deepseek-chat")
    available = capability.available_input_tokens(system_tokens=500, schema_tokens=600)

    assert capability.context_window > capability.max_output_tokens
    assert available < capability.context_window
    assert estimate_tokens("线性代数" * 100) > len("线性代数" * 100)


def test_oversized_chapter_is_semantically_split_to_token_budget() -> None:
    text = "\n\n".join(
        f"## Section {index}\n" + ("向量空间、基与维数。 " * 300)
        for index in range(20)
    )
    chunks = split_chapter(ChapterSource("Long chapter", text), token_budget=900)

    assert len(chunks) > 2
    assert all(estimate_tokens(chunk.text) <= 900 for chunk in chunks)
    assert all(chunk.section_title for chunk in chunks)


def test_chapter_sources_ignore_table_of_contents_clusters_and_running_duplicates() -> None:
    parsed = parsed_material()
    parsed.pages = [
        ParsedPage(page_number=index, text=f"第 {index} 页内容", source="text_extract", source_anchor=f"linear.pdf, p. {index}")
        for index in range(1, 13)
    ]
    parsed.raw_text = "\n".join(page.text for page in parsed.pages)
    parsed.document_structure = DocumentStructure(
        document_id="doc-toc",
        filename="linear.pdf",
        file_type=".pdf",
        title="Linear algebra",
        page_count=12,
        sections=[
            DocumentSection(section_id="toc-1", title="第一章 行列式", page_start=1, page_end=1, source_anchor="linear.pdf, p. 1"),
            DocumentSection(section_id="toc-2", title="第二章 矩阵", page_start=1, page_end=1, source_anchor="linear.pdf, p. 1"),
            DocumentSection(section_id="toc-3", title="第三章 向量", page_start=1, page_end=1, source_anchor="linear.pdf, p. 1"),
            DocumentSection(section_id="c1", title="第一章 行列式", page_start=3, page_end=5, source_anchor="linear.pdf, p. 3"),
            DocumentSection(section_id="c1-header", title="第一章 行列式", page_start=4, page_end=4, source_anchor="linear.pdf, p. 4"),
            DocumentSection(section_id="c2", title="第二章 矩阵", page_start=6, page_end=8, source_anchor="linear.pdf, p. 6"),
            DocumentSection(section_id="c3", title="第三章 向量", page_start=9, page_end=12, source_anchor="linear.pdf, p. 9"),
        ],
    )

    chapters = build_chapter_sources([parsed])

    assert [chapter.title for chapter in chapters] == ["第一章 行列式", "第二章 矩阵", "第三章 向量"]
    assert "linear.pdf, p. 3" in chapters[0].text
    assert "linear.pdf, p. 1" not in chapters[0].text


def test_finish_reason_length_is_explicit_truncation(monkeypatch) -> None:
    provider = get_llm_provider("openai")

    def fake_post(*args, **kwargs):
        return httpx.Response(
            200,
            json={
                "choices": [{"finish_reason": "length", "message": {"content": "{\"name\":\"cut"}}],
                "usage": {"prompt_tokens": 100, "completion_tokens": 50},
            },
        )

    monkeypatch.setattr("app.services.llm_providers.openai_compatible.httpx.post", fake_post)
    with pytest.raises(LLMProviderError) as captured:
        provider.post_chat_completions(
            "https://example.test/v1/chat/completions",
            "not-logged",
            {"model": "gpt-4o-mini", "messages": []},
            timeout=1,
        )
    assert captured.value.error.code == "LLM_OUTPUT_TRUNCATED"


def test_incomplete_json_at_output_cap_is_explicit_truncation(monkeypatch) -> None:
    provider = get_llm_provider("openai")
    monkeypatch.setattr(
        "app.services.llm_providers.openai_compatible.httpx.post",
        lambda *args, **kwargs: httpx.Response(
            200,
            json={
                "choices": [{"finish_reason": None, "message": {"content": "{\"name\":\"cut\""}}],
                "usage": {"prompt_tokens": 100, "completion_tokens": 512},
            },
        ),
    )

    with pytest.raises(LLMProviderError) as captured:
        provider.generate_study_unit(
            "Chapter",
            "source",
            LLMConfig(provider="openai", model="gpt-4o-mini", api_key="secret", enabled=True),
            max_output_tokens=512,
        )

    assert captured.value.error.code == "LLM_OUTPUT_TRUNCATED"


def test_balanced_malformed_json_is_parse_error_not_truncation(monkeypatch) -> None:
    provider = get_llm_provider("openai")
    monkeypatch.setattr(
        "app.services.llm_providers.openai_compatible.httpx.post",
        lambda *args, **kwargs: httpx.Response(
            200,
            json={
                "choices": [{"finish_reason": "stop", "message": {"content": "{\"name\": ]}"}}],
                "usage": {"prompt_tokens": 100, "completion_tokens": 20},
            },
        ),
    )

    with pytest.raises(LLMProviderError) as captured:
        provider.generate_study_unit(
            "Chapter",
            "source",
            LLMConfig(provider="openai", model="gpt-4o-mini", api_key="secret", enabled=True),
            max_output_tokens=512,
        )

    assert captured.value.error.code == "LLM_RESPONSE_PARSE_ERROR"


def test_truncated_study_unit_retries_then_completes(monkeypatch) -> None:
    provider = get_llm_provider("openai")
    calls = 0
    seen_lengths: list[int] = []

    def fake_generate(_title, chunk_text, *_args, **_kwargs):
        nonlocal calls
        calls += 1
        seen_lengths.append(len(chunk_text))
        if calls <= 2:
            raise LLMProviderError(
                LLMErrorInfo(
                    code="LLM_OUTPUT_TRUNCATED",
                    message="truncated",
                    suggestion="retry",
                    provider="OpenAI",
                    model="gpt-4o-mini",
                )
            )
        return unit()

    monkeypatch.setattr(provider, "generate_study_unit", fake_generate)
    parsed = parsed_material()
    result, stats = generate_hierarchical_report(
        [parsed],
        generate_review_report(parsed.raw_text),
        LLMConfig(provider="openai", model="gpt-4o-mini", api_key="secret", enabled=True),
        job_id="job-truncate",
    )

    assert result.llm_status == "success"
    assert stats.llm_calls == 4
    assert stats.retry_count == 2
    assert result.report.study_units[0].name == "Vector spaces"
    assert seen_lengths[1] == seen_lengths[0]
    assert seen_lengths[2] < seen_lengths[0]
    assert seen_lengths[3] < seen_lengths[0]


def test_split_recovery_resumes_completed_child_checkpoint(monkeypatch) -> None:
    provider = get_llm_provider("openai")
    calls: list[str] = []

    def fake_generate(_title, chunk_text, *_args, **_kwargs):
        calls.append(chunk_text)
        if len(calls) <= 2:
            raise LLMProviderError(
                LLMErrorInfo(
                    code="LLM_CONTEXT_EXCEEDED",
                    message="too long",
                    suggestion="split",
                    provider="OpenAI",
                    model="gpt-4o-mini",
                )
            )
        return unit("Vector spaces")

    monkeypatch.setattr(provider, "generate_study_unit", fake_generate)
    parsed = parsed_material("第一段向量空间定义与性质。\n\n第二段线性方程组求解方法。")
    result, stats = generate_hierarchical_report(
        [parsed],
        generate_review_report(parsed.raw_text),
        LLMConfig(provider="openai", model="gpt-4o-mini", api_key="secret", enabled=True),
        job_id="job-split-resume",
        checkpoints={
            "chapter-1:chunk-1.1": {
                "status": "completed",
                "content": unit("Saved first half").model_dump(mode="json"),
            }
        },
    )

    assert result.llm_status == "success"
    assert stats.llm_calls == 3
    assert stats.retry_count == 2
    assert stats.resumed_checkpoints == 1
    assert result.report.study_units


@pytest.mark.parametrize(
    ("status", "body", "expected"),
    [
        (429, {"error": {"message": "rate limit"}}, "RATE_LIMITED"),
        (429, {"error": {"message": "insufficient quota"}}, "PROVIDER_QUOTA_EXCEEDED"),
        (503, {"error": {"message": "upstream unavailable"}}, "LLM_PROVIDER_ERROR"),
    ],
)
def test_provider_http_errors_have_specific_codes(monkeypatch, status: int, body: dict, expected: str) -> None:
    provider = get_llm_provider("openai")
    monkeypatch.setattr(
        "app.services.llm_providers.openai_compatible.httpx.post",
        lambda *args, **kwargs: httpx.Response(status, json=body),
    )

    with pytest.raises(LLMProviderError) as captured:
        provider.post_chat_completions(
            "https://example.test/v1/chat/completions",
            "secret",
            {"model": "gpt-4o-mini", "messages": []},
            timeout=1,
        )

    assert captured.value.error.code == expected


def test_generation_diagnostics_redacts_source_and_secrets(caplog) -> None:
    caplog.set_level("INFO", logger="recallforge.generation")
    with diagnostic_context(job_id="job-safe", current_chunk=2):
        diagnostic(
            "test_event",
            api_key="secret-key",
            authorization="Bearer secret",
            chunk_text="教材完整正文",
            returned_character_count=321,
        )

    assert "job-safe" in caplog.text
    assert "returned_character_count" in caplog.text
    assert "secret-key" not in caplog.text
    assert "教材完整正文" not in caplog.text


@pytest.mark.parametrize(
    ("source_code", "expected_code"),
    [
        ("TIMEOUT", "LLM_TIMEOUT"),
        ("RESPONSE_PARSE_ERROR", "LLM_RESPONSE_PARSE_ERROR"),
        ("RATE_LIMITED", "PROVIDER_RATE_LIMIT"),
        ("LLM_PROVIDER_ERROR", "LLM_PROVIDER_ERROR"),
    ],
)
def test_retryable_provider_failures_are_classified_and_keep_partial_material(
    monkeypatch, source_code: str, expected_code: str
) -> None:
    provider = get_llm_provider("openai")

    def fail(*args, **kwargs):
        raise LLMProviderError(
            LLMErrorInfo(
                code=source_code,
                message="provider failed",
                suggestion="retry",
                provider="OpenAI",
                model="gpt-4o-mini",
            ),
            http_status=429 if source_code == "RATE_LIMITED" else 503,
        )

    monkeypatch.setattr(provider, "generate_study_unit", fail)
    parsed = parsed_material()
    result, stats = generate_hierarchical_report(
        [parsed],
        generate_review_report(parsed.raw_text),
        LLMConfig(provider="openai", model="gpt-4o-mini", api_key="secret", enabled=True),
        job_id=f"job-{source_code}",
    )

    assert result.llm_status == "failed"
    assert result.llm_error and result.llm_error.code == expected_code
    assert stats.retry_count == 2
    assert result.report.study_units


def test_completed_checkpoint_is_resumed_without_llm_call(monkeypatch) -> None:
    provider = get_llm_provider("openai")
    monkeypatch.setattr(
        provider,
        "generate_study_unit",
        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("LLM must not run")),
    )
    parsed = parsed_material()
    result, stats = generate_hierarchical_report(
        [parsed],
        generate_review_report(parsed.raw_text),
        LLMConfig(provider="openai", model="gpt-4o-mini", api_key="secret", enabled=True),
        job_id="job-resume",
        checkpoints={
            "chapter-1:chunk-1": {
                "status": "completed",
                "content": unit("Saved unit").model_dump(mode="json"),
            }
        },
    )

    assert stats.resumed_checkpoints == 1
    assert stats.llm_calls == 0
    assert result.report.study_units[0].name == "Saved unit"


def test_generation_store_persists_result_parse_and_checkpoint(tmp_path: Path, monkeypatch) -> None:
    database = tmp_path / "jobs.sqlite3"
    monkeypatch.setenv("RECALLFORGE_JOB_DB", str(database))
    payload = {
        "files": ["saved.pdf"],
        "project_id": "project-1",
        "llm_config": {"provider": "openai", "model": "gpt-4o-mini", "api_key": "must-not-persist", "enabled": True},
    }
    job_id, resumed = generation_store.create_or_resume("job-store", payload)
    generation_store.save_parsed(job_id, [parsed_material().model_dump(mode="json")])
    generation_store.save_checkpoint(job_id, "chapter-1:chunk-1", 1, 1, "completed", unit().model_dump(mode="json"), None, None)
    generation_store.update_job(job_id, status="completed", progress=100, message="done", result={"ok": True})

    stored = generation_store.get_job(job_id)
    assert resumed is False
    assert stored and stored["status"] == "completed" and stored["result"] == {"ok": True}
    assert generation_store.load_parsed(job_id)
    assert generation_store.load_checkpoints(job_id)["chapter-1:chunk-1"]["status"] == "completed"
    assert "must-not-persist" not in database.read_bytes().decode("utf-8", errors="ignore")

    generation_store.update_job(
        job_id,
        status="retryable_failed",
        progress=100,
        message="retry",
        error_code="LLM_TIMEOUT",
        retryable=True,
    )
    resumed_id, resumed = generation_store.create_or_resume(
        "must-not-create",
        {**payload, "llm_config": {**payload["llm_config"], "api_key": "new-secret"}},
    )
    resumed_job = generation_store.get_job(resumed_id)
    assert resumed is True and resumed_id == job_id
    assert resumed_job and resumed_job["status"] == "retrying"
    assert "new-secret" not in database.read_bytes().decode("utf-8", errors="ignore")


def test_generation_store_resumes_unscoped_request(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("RECALLFORGE_JOB_DB", str(tmp_path / "jobs.sqlite3"))
    payload = {"files": ["saved.pdf"], "llm_config": {"enabled": False}}
    job_id, _ = generation_store.create_or_resume("unscoped-job", payload)
    generation_store.update_job(
        job_id,
        status="retryable_failed",
        progress=100,
        message="retry",
        error_code="LLM_TIMEOUT",
        retryable=True,
    )

    resumed_id, resumed = generation_store.create_or_resume("new-job", payload)

    assert resumed is True
    assert resumed_id == job_id
    assert generation_store.get_job(job_id)["status"] == "retrying"


def test_generation_success_survives_export_failure(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("RECALLFORGE_JOB_DB", str(tmp_path / "jobs.sqlite3"))
    monkeypatch.setattr(settings, "output_dir", tmp_path / "outputs")
    report = ReviewReport(title="Persisted", summary="Body", chapters=[], markdown="# Persisted\n\nBody")

    def fake_build(request, progress_callback=None, **kwargs):
        return GenerateReviewResponse(
            review_report=report,
            markdown=report.markdown,
            export_format="md",
            llm_status="disabled",
        )

    monkeypatch.setattr("app.routers.generate_review_jobs_v31.build_generate_review_response", fake_build)
    client = TestClient(app)
    created = client.post("/api/review/jobs", json={"files": ["demo.pdf"], "project_id": "project-export"}).json()
    job_id = created["job_id"]
    for _ in range(100):
        job = client.get(f"/api/review/jobs/{job_id}").json()
        if job["status"] == "completed":
            break
        time.sleep(0.01)
    else:
        raise AssertionError("job did not complete")

    monkeypatch.setattr(
        "app.routers.generate_review_jobs_v31.export_review_report",
        lambda *args, **kwargs: (_ for _ in ()).throw(ExportError("renderer failed")),
    )
    failed_export = client.get(f"/api/review/jobs/{job_id}/download/md")

    assert failed_export.status_code == 500
    assert "EXPORT_RENDER_ERROR" in failed_export.json()["detail"]
    persisted = client.get(f"/api/review/jobs/{job_id}").json()
    assert persisted["status"] == "completed"
    assert persisted["result"]["markdown"].startswith("# Persisted")
    assert not list((tmp_path / "outputs").glob("recallforge-export-*"))


def test_large_study_material_exports_all_supported_formats(tmp_path: Path) -> None:
    body = """
# 线性代数复习资料

## 向量空间 Vector Space

- 中文与 English
- inline math: $Ax=b$
- 特殊字符：∑ ∫ √ ≤ ≥ ± ∞

| 概念 | 定义 |
| --- | --- |
| 基 | 线性无关生成组 |

$$A^{-1}A=I$$
""".strip()
    sentinel = "FINAL-SENTINEL-RECALLFORGE-EXPORT"
    markdown = "\n\n".join(f"{body}\n\n### 长章节 {index}\n正文 {index}。" for index in range(80)) + f"\n\n## 完整性校验\n\n{sentinel}"
    report = ReviewReport(
        title="线性代数复习资料",
        summary="完整正文",
        chapters=[],
        markdown=markdown,
        anki_cards=[AnkiCard(front="向量空间是什么？", back="满足线性运算封闭性的集合。", tags="线性代数")],
    )

    md = export_review_report(report, markdown, tmp_path, "large", "md")
    docx = export_review_report(report, markdown, tmp_path, "large", "docx")
    pdf = export_review_report(report, markdown, tmp_path, "large", "pdf")
    anki = export_anki_csv(report, tmp_path, "large-anki.csv")

    assert md.stat().st_size > 0 and md.read_text(encoding="utf-8").endswith(sentinel)
    reopened_docx = Document(docx)
    assert any("线性代数复习资料" in paragraph.text for paragraph in reopened_docx.paragraphs)
    assert sum(len(paragraph.text) for paragraph in reopened_docx.paragraphs) > 1_000
    assert any(sentinel in paragraph.text for paragraph in reopened_docx.paragraphs)
    reopened_pdf = PdfReader(str(pdf))
    assert len(reopened_pdf.pages) > 0
    assert pdf.stat().st_size > 0
    assert sentinel in "".join(page.extract_text() or "" for page in reopened_pdf.pages)
    assert "向量空间是什么" in anki.read_text(encoding="utf-8-sig")


def test_sqlite_persists_large_canonical_material_without_truncation(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("RECALLFORGE_JOB_DB", str(tmp_path / "jobs.sqlite3"))
    payload = {"files": ["long.pdf"], "project_id": "long-project", "llm_config": {"enabled": False}}
    job_id, _ = generation_store.create_or_resume("job-large-persist", payload)
    sentinel = "FINAL-SENTINEL-PERSIST"
    markdown = ("章节正文与公式 Ax=b。\n" * 80_000) + sentinel
    result = {
        "review_report": ReviewReport(title="Long", summary="Body", markdown=markdown).model_dump(mode="json"),
        "markdown": markdown,
        "export_format": "md",
    }

    generation_store.update_job(job_id, status="completed", progress=100, message="done", result=result)
    stored = generation_store.get_job(job_id)

    assert stored and stored["result"]["markdown"].endswith(sentinel)
    assert len(stored["result"]["markdown"]) == len(markdown)
